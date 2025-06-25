import os
import subprocess
from pathlib import Path
from fastapi import FastAPI, Form
from fastapi.middleware.cors import CORSMiddleware
from docx import Document as DocxDocument
from langchain.prompts import PromptTemplate
from langchain.chains.combine_documents import create_stuff_documents_chain
from langchain_community.vectorstores import Chroma
from langchain_core.runnables import RunnableLambda, RunnablePassthrough
from langchain_core.documents import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_google_genai.chat_models import ChatGoogleGenerativeAI
from langchain_google_genai import ChatGoogleGenerativeAI

# === Configuration ===
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY", "apikey")

# === FastAPI App Setup ===
app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# === Constants ===
DOCS_DIR = "/home/diwan/Downloads/smart_report"
DB_DIR = "medical_db"

# === Prompt Template ===
prompt = PromptTemplate.from_template("""
You are a licensed medical assistant Answer the following question based on your medical knowledge and on the provided medical documents if given any .:\n\n"
Context: {context}
Question: {query}
Answer:""")

# === Components ===
llm = ChatGoogleGenerativeAI(
    model="gemini-1.5-flash",
    temperature=0.6,
    google_api_key="AIzaSyAilOaRoHx5sO7seufB2SMuX7tusoAbh3I"
)
splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
embeddings = HuggingFaceEmbeddings(model_name="intfloat/multilingual-e5-base")

# === Globals for Lazy Initialization ===
vectorstore = None
retriever = None


# === Helper Functions ===
def format_docs(docs):
    """Robust document formatter that handles various input types"""
    print(f"🔧 format_docs input type: {type(docs)}")

    if not docs:
        print("🔧 Empty docs received")
        return ""

    # Handle single Document object
    if hasattr(docs, 'page_content'):
        print("🔧 Single Document object")
        return docs.page_content

    # Handle string input
    if isinstance(docs, str):
        print("🔧 String input")
        return docs

    # Handle list of documents
    if isinstance(docs, list):
        print(f"🔧 List of {len(docs)} items")
        formatted_docs = []
        for i, doc in enumerate(docs):
            print(f"🔧 Processing item {i}: {type(doc)}")
            if hasattr(doc, 'page_content'):
                formatted_docs.append(doc.page_content)
            elif isinstance(doc, str):
                formatted_docs.append(doc)
            else:
                print(f"🔧 Unknown doc type: {type(doc)}")
                formatted_docs.append(str(doc))

        result = "\n\n".join(formatted_docs)
        print(f"🔧 Formatted result length: {len(result)}")
        return result

    # Fallback for unknown types
    print(f"🔧 Unknown input type, converting to string: {type(docs)}")
    return str(docs)


def get_context(query):
    """Retrieve and format context for a query"""
    global retriever
    if not retriever:
        print("❌ Retriever not initialized")
        return "No context available"

    try:
        print(f"🔍 Retrieving docs for query: {query}")
        docs = retriever.invoke(query)
        print(f"🔍 Retrieved {len(docs)} documents")

        # Debug each retrieved document
        for i, doc in enumerate(docs):
            print(f"🔍 Doc {i} type: {type(doc)}")
            if hasattr(doc, 'page_content'):
                print(f"🔍 Doc {i} content length: {len(doc.page_content)}")

        return format_docs(docs)
    except Exception as e:
        print(f"❌ Error in get_context: {e}")
        import traceback
        traceback.print_exc()
        return "Error retrieving context"


# === Utilities ===
def convert_doc_to_docx(path: Path):
    if path.suffix.lower() == ".doc":
        try:
            subprocess.run([
                "libreoffice", "--headless", "--convert-to", "docx", str(path), "--outdir", str(path.parent)
            ], check=True)
            print(f"✅ Converted: {path.name}")
            return path.with_suffix(".docx")
        except subprocess.CalledProcessError as e:
            print(f"❌ Failed to convert {path.name}: {e}")
            return None
    return path


def load_docx_text(path: Path) -> str:
    try:
        doc = DocxDocument(path)
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        print(f"❌ Failed to read {path.name}: {e}")
        return ""


def ingest_documents():
    print("📄 ingest_documents() called")
    docs = []
    os.makedirs(DB_DIR, exist_ok=True)

    for file_path in Path(DOCS_DIR).rglob("*"):
        if file_path.is_file() and file_path.suffix.lower() in [".doc", ".docx"]:
            try:
                docx_path = convert_doc_to_docx(file_path)
                if docx_path and docx_path.suffix.lower() == ".docx":
                    text = load_docx_text(docx_path)
                    if text.strip():
                        docs.append(Document(page_content=text, metadata={"source": docx_path.name}))
                        print(f"✅ Loaded: {docx_path.name}")
                    else:
                        print(f"⚠️ Empty content in: {docx_path.name}")
            except Exception as e:
                print(f"❌ Error processing {file_path.name}: {e}")

    if not docs:
        print("⚠️ No documents loaded.")
        return None

    print(f"📄 Loaded {len(docs)} documents")
    chunks = splitter.split_documents(docs)
    print(f"📄 Created {len(chunks)} chunks")

    # Verify chunks are proper Document objects
    for i, chunk in enumerate(chunks[:3]):  # Check first 3 chunks
        print(f"📄 Chunk {i} type: {type(chunk)}")
        if hasattr(chunk, 'page_content'):
            print(f"📄 Chunk {i} content length: {len(chunk.page_content)}")

    print("⚙️ Creating Chroma vectorstore...")
    try:
        vectorstore = Chroma.from_documents(
            documents=chunks,
            embedding=embeddings,
            persist_directory=DB_DIR
        )
        print("💾 Persisting Chroma DB to disk...")
        vectorstore.persist()
        print("✅ Done persisting vectorstore.")
        return vectorstore
    except Exception as e:
        print(f"❌ Error creating vectorstore: {e}")
        import traceback
        traceback.print_exc()
        return None


def ensure_vectorstore():
    global vectorstore, retriever
    if vectorstore is None:
        if os.path.exists(DB_DIR):
            print("📦 Loading existing Chroma DB...")
            try:
                vectorstore = Chroma(persist_directory=DB_DIR, embedding_function=embeddings)
                print("✅ Successfully loaded existing vectorstore")
            except Exception as e:
                print(f"❌ Failed to load existing DB: {e}")
                print("📄 Creating new vectorstore...")
                vectorstore = ingest_documents()
        else:
            print("📄 DB not found. Ingesting documents...")
            vectorstore = ingest_documents()

        if vectorstore:
            print("⚙️ Setting up retriever...")
            retriever = vectorstore.as_retriever(search_kwargs={"k": 5})
            print("✅ Retriever initialized")
        else:
            print("❌ Failed to initialize vectorstore")


# === API Endpoints ===
@app.post("/question")
async def ask_question(query: str = Form(...)):
    try:
        print(f"🔍 Query received: {query}")
        ensure_vectorstore()

        # Default fallback context handling
        context = ""
        sources = []

        if retriever:
            print("⚙️ Trying to retrieve context...")
            context = get_context(query)

        # Check if we should skip the document context
        use_context = context and context != "Error retrieving context" and context.strip()

        # Use template with or without context
        if use_context:
            print(f"✅ Using retrieved context (length={len(context)})")
            formatted_prompt = prompt.format(context=context, query=query)
        else:
            print("⚠️ No relevant context found. Falling back to direct LLM invocation.")
            formatted_prompt = f"You are a medical assistant. Answer this question:\n\nQuestion: {query}\nAnswer:"

        # Get response from LLM
        print("⚙️ Invoking LLM...")
        result = llm.invoke(formatted_prompt)

        # Extract answer
        answer = result.content if hasattr(result, 'content') else str(result)

        # If context was used, provide source documents
        if use_context:
            print("🔍 Getting sources from vectorstore...")
            docs = vectorstore.similarity_search(query, k=5)
            sources = [doc.metadata.get("source", "unknown") for doc in docs if hasattr(doc, 'metadata')]

        print(f"✅ Answer generated. Sources: {sources}")
        return {"answer": answer, "sources": sources}

    except Exception as e:
        print(f"❌ Error in ask_question: {str(e)}")
        import traceback
        traceback.print_exc()
        return {"error": f"Internal error: {str(e)}"}



@app.post("/refresh")
def refresh_docs():
    global vectorstore, retriever
    try:
        print("🔄 Refreshing documents...")
        vectorstore = ingest_documents()

        if vectorstore:
            retriever = vectorstore.as_retriever(search_kwargs={"k": 5})
            return {"status": "✅ Documents re-ingested."}
        return {"error": "⚠️ No documents to ingest."}
    except Exception as e:
        print(f"❌ Error in refresh_docs: {str(e)}")
        return {"error": str(e)}


@app.get("/health")
def health_check():
    """Health check endpoint"""
    return {"status": "✅ API is running"}


@app.get("/debug")
def debug_info():
    """Debug endpoint to check vectorstore status"""
    global vectorstore, retriever
    ensure_vectorstore()

    info = {
        "vectorstore_initialized": vectorstore is not None,
        "retriever_initialized": retriever is not None,
        "db_dir_exists": os.path.exists(DB_DIR),
        "docs_dir_exists": os.path.exists(DOCS_DIR)
    }

    if vectorstore:
        try:
            # Test retrieval
            test_docs = vectorstore.similarity_search("test", k=1)
            info["test_retrieval_count"] = len(test_docs)
            if test_docs:
                info["test_doc_type"] = str(type(test_docs[0]))
        except Exception as e:
            info["test_retrieval_error"] = str(e)

    return info


@app.get("/")
def root():
    """Root endpoint with basic info"""
    return {"message": "Medical RAG Assistant API", "endpoints": ["/question", "/refresh", "/health", "/debug"]}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="11.11.11.222", port=8060)